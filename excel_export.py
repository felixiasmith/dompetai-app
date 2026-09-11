import os
from datetime import datetime, timezone
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from motor.motor_asyncio import AsyncIOMotorClient
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/api/v1/export", tags=["Export Engine"])

# Koneksi Database Helper
async def get_db_client():
    mongo_uri = os.getenv("MONGO_URI", "mongodb://localhost:27017")
    return AsyncIOMotorClient(mongo_uri)

@router.get("/excel/{user_id}")
async def export_financial_excel(user_id: str, db_client: AsyncIOMotorClient = Depends(get_db_client)):
    db = db_client["dompet_ai_db"]
    
    # 1. Ambil seluruh data dari MongoDB secara Asynchronous
    transactions = await db.transactions.find({"user_id": user_id}).sort("timestamp_utc", -1).to_list(length=5000)
    wallets = await db.wallets.find({"user_id": user_id}).to_list(length=100)
    debts = await db.debts_receivables.find({"user_id": user_id}).to_list(length=500)

    # 2. Inisialisasi Workbook OpenPyXL
    wb = openpyxl.Workbook()
    
    # Atur Style Global untuk Tampilan Premium (Emerald & Slate theme)
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=11)
    fill_header = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid") # Emerald Green
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    
    border_thin = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    # ----------------------------------------------------
    # SHEET 1: RINGKASAN DOMPET (WALLETS)
    # ----------------------------------------------------
    ws_wallets = wb.active
    ws_wallets.title = "Status Dompet"
    ws_wallets.views.sheetView[0].showGridLines = True
    
    wallet_headers = ["Nama Dompet", "Tipe", "Saldo Saat Ini", "Mata Uang"]
    ws_wallets.append(wallet_headers)
    
    for w in wallets:
        ws_wallets.append([w["name"], w["type"], w["balance"], w["currency"]])
        
    # Formatting Sheet Dompet
    for row in ws_wallets.iter_rows(min_row=1, max_row=ws_wallets.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.font = font_data
            cell.border = cell_border
            if cell.row == 1:
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = align_center
            elif cell.column == 3:
                cell.number_format = '#,##0'
                cell.alignment = align_right

    # ----------------------------------------------------
    # SHEET 2: DAFTAR TRANSAKSI (LEDGER)
    # ----------------------------------------------------
    ws_tx = wb.create_sheet(title="Riwayat Transaksi")
    ws_tx.views.sheetView[0].showGridLines = True
    
    tx_headers = ["Tanggal (UTC)", "Deskripsi Transaksi", "Nominal", "Tipe", "Kategori 50/30/20", "Catatan Asli AI"]
    ws_tx.append(tx_headers)
    
    for t in transactions:
        # Menghapus timezone awareness agar kompatibel sempurna saat ditulis ke Excel
        dt_naive = t["timestamp_utc"].replace(tzinfo=None) if isinstance(t["timestamp_utc"], datetime) else t["timestamp_utc"]
        ws_tx.append([
            dt_naive,
            t["description"],
            t["amount"],
            t["transaction_type"],
            t["category"],
            t.get("raw_nlp_text", "")
        ])
        
    # Formatting Sheet Transaksi
    for row in ws_tx.iter_rows(min_row=1, max_row=ws_tx.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.font = font_data
            cell.border = cell_border
            if cell.row == 1:
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = align_center
            else:
                if cell.column == 1:
                    cell.number_format = 'yyyy-mm-dd hh:mm:ss'
                    cell.alignment = align_center
                elif cell.column == 3:
                    cell.number_format = '#,##0'
                    cell.alignment = align_right
                elif cell.column in:
                    cell.alignment = align_center

    # ----------------------------------------------------
    # SHEET 3: UTANG & PIUTANG (DEBTS)
    # ----------------------------------------------------
    ws_debts = wb.create_sheet(title="Utang Piutang (IOU)")
    ws_debts.views.sheetView[0].showGridLines = True
    
    debt_headers = ["Jenis", "Nama Orang", "Total Pinjaman", "Sisa Belum Lunas", "Status", "Keterangan"]
    ws_debts.append(debt_headers)
    
    for d in debts:
        ws_debts.append([
            d["type"],
            d["person_name"],
            d["amount"],
            d["remaining_amount"],
            d["status"],
            d["description"]
        ])
        
    # Formatting Sheet Utang
    for row in ws_debts.iter_rows(min_row=1, max_row=ws_debts.max_row, min_col=1, max_col=6):
        for cell in row:
            cell.font = font_data
            cell.border = cell_border
            if cell.row == 1:
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = align_center
            else:
                if cell.column in:
                    cell.number_format = '#,##0'
                    cell.alignment = align_right
                elif cell.column in:
                    cell.alignment = align_center

    # ----------------------------------------------------
    # OTOMATISASI UKURAN KOLOM (Auto-fit Width)
    # ----------------------------------------------------
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # 3. Simpan Workbook ke dalam Memory Stream (Tanpa Menulis File ke Disk Server)
    file_stream = BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    
    filename = f"Laporan_Keuangan_DompetAI_{datetime.now().strftime('%Y%m%d')}.xlsx"
    
    # 4. Stream File Langsung ke Web Browser Pengguna
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
